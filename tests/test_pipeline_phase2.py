"""Full daily run with auto-apply and outreach on: job boards (fixtures) -> scoring (fake Claude) -> apply (fake
browser serving the real captured forms) -> outreach (Apollo + Gmail mocked over HTTP) -> workbook -> summary."""

import json
import logging
from pathlib import Path

import pytest
import respx
import yaml
from openpyxl import load_workbook

from jobagent.config import load_config
from jobagent.db import Store
from jobagent.http import HttpClient
from jobagent.logging_setup import setup_logging
from jobagent.outreach.apollo import ApolloClient
from jobagent.pipeline import Pipeline, RunOptions

from .conftest import FakeLLM
from .test_apply_engine import FakeDriver
from .test_outreach import mock_apollo, mock_gmail
from .test_pipeline import mock_boards

FORMS = Path(__file__).parent / "fixtures" / "forms"


class AnyUrlDriver(FakeDriver):
    """Serves the real Lever form for Lever URLs and the real Greenhouse form for everything else."""

    async def extract(self):
        name = "lever_synthego" if "lever" in self.url else "greenhouse_xaira"
        return json.loads((FORMS / f"{name}.json").read_text())


@pytest.fixture
def cfg(project):
    data = yaml.safe_load((project / "kb" / "answers.example.yaml").read_text())
    data["identity"].update(legal_first_name="Riya", legal_last_name="Sancheti", email="riya@example.com",
                            phone="+1 555 0100", linkedin_url="https://linkedin.com/in/example")
    (project / "kb" / "answers.yaml").write_text(yaml.safe_dump(data))
    for name in ("resume_business.pdf", "resume_research.pdf"):
        (project / "kb" / "resumes" / name).write_bytes(b"%PDF-1.4\n")
    (project / "credentials").mkdir()
    (project / "credentials" / "gmail_token.json").write_text(json.dumps({
        "token": "t", "refresh_token": "r", "client_id": "c", "client_secret": "s",
        "token_uri": "https://oauth2.googleapis.com/token", "expiry": "2999-01-01T00:00:00Z"}))
    c = load_config(root=project)
    c.data["notify"]["summary_email"] = True
    c.data["apply"]["pause_seconds"] = [0, 0]
    c.data["outreach"]["mode"] = "draft"
    setup_logging(c.path("logs_dir"), "DEBUG", "WARNING", 30, True)
    yield c
    logging.getLogger().handlers.clear()


@respx.mock
async def test_full_day_with_apply_and_outreach(cfg):
    mock_boards()
    mock_apollo()
    drafts, send, _ = mock_gmail()
    store = Store(cfg.path("db"))
    llm = FakeLLM()
    driver = AnyUrlDriver({}, outcome="confirmed")
    async with HttpClient(max_retries=0) as http:
        p = Pipeline(cfg, store, http, lambda: llm,
                     RunOptions(apply=True, outreach=True, dry_run=False), driver_factory=lambda: driver,
                     apollo=ApolloClient(http, "test-key"))
        stats = await p.run()

    assert stats["status"] == "ok", stats
    applied = store.query("SELECT * FROM jobs WHERE apply_status='applied'")
    assert len(applied) >= 3
    assert any(c[0] == "upload" and "cover_letter" in c[2] for c in driver.calls)      # Greenhouse required one
    purposes = {r["purpose"] for r in store.query("SELECT DISTINCT purpose FROM llm_usage")}
    assert {"score", "answers", "cover_letter", "email"} <= purposes
    assert stats["outreach"]["contacts_found"] == 1 and stats["outreach"]["company_cooldown"] == 1
    assert send.call_count == 1                                                         # only the daily summary
    wb = load_workbook(cfg.path("workbook"))
    applied_ws, contacts_ws = wb["Applied"], wb["Contacts"]
    assert applied_ws.max_row - 1 == len(applied)
    rows = [[c.value for c in r] for r in contacts_ws.iter_rows(min_row=2)]
    dana = [r for r in rows if r[5] == "dana@acmebio.com"]
    assert len(dana) == 1 and dana[0][7] == "drafted"                     # one email to Dana, not one per job
    assert any(r[3] == "(find manually)" for r in rows)                  # ChipCo: no domain found
    assert drafts.call_count == 1
    run = store.query("SELECT * FROM runs")[0]
    assert run["cost_usd"] > 0 and json.loads(run["stats"])["apply"]["applied"] >= 3
    store.close()


@respx.mock
async def test_outreach_only_run(cfg):
    mock_apollo()
    mock_gmail()
    store = Store(cfg.path("db"))
    async with HttpClient(max_retries=0) as http:
        stats = await Pipeline(cfg, store, http, lambda: FakeLLM(), RunOptions(only_outreach=True, outreach=True)).run()
    assert stats["status"] == "ok" and "outreach" in stats and "jobs_seen" not in stats
    store.close()
