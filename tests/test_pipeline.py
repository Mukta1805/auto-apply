"""End-to-end: mocked job boards + fake Claude -> database -> workbook -> Riya's edits survive a rebuild."""

import json
import logging

import httpx
import pytest
import respx
from openpyxl import load_workbook

from jobagent.config import load_config
from jobagent.db import Store
from jobagent.export import build_workbook
from jobagent.http import HttpClient
from jobagent.logging_setup import mask_secrets, setup_logging
from jobagent.pipeline import Pipeline, RunLock, RunOptions

from .conftest import FakeLLM, load_fixture

SR = "https://api.smartrecruiters.com/v1/companies/ConsultCo/postings"


def mock_boards():
    respx.get("https://boards-api.greenhouse.io/v1/boards/acmebio/jobs").mock(
        return_value=httpx.Response(200, json=load_fixture("greenhouse_jobs.json"), headers={"ETag": '"v1"'}))
    respx.get("https://api.lever.co/v0/postings/chipco").mock(
        return_value=httpx.Response(200, json=load_fixture("lever_postings.json")))
    respx.get("https://api.ashbyhq.com/posting-api/job-board/aibio").mock(
        return_value=httpx.Response(500))                           # one board down: run must continue
    respx.get("https://apply.workable.com/api/v1/widget/accounts/medtechco").mock(
        return_value=httpx.Response(200, json=load_fixture("workable_account.json")))
    respx.get(SR).mock(return_value=httpx.Response(200, json=load_fixture("smartrecruiters_postings.json")))
    respx.get(f"{SR}/744000011").mock(return_value=httpx.Response(200, json=load_fixture("smartrecruiters_detail.json")))


@pytest.fixture
def cfg(project):
    c = load_config(root=project)
    setup_logging(c.path("logs_dir"), "DEBUG", "WARNING", 30, True)
    yield c
    logging.getLogger().handlers.clear()


@respx.mock
async def test_full_run(cfg, monkeypatch):
    mock_boards()
    store = Store(cfg.path("db"))
    llm = FakeLLM()
    async with HttpClient(max_retries=1) as http:
        monkeypatch.setattr(http, "_retry_delay", lambda resp, attempt: 0)
        import jobagent.http as h
        monkeypatch.setattr(h.random, "random", lambda: 0)
        stats = await Pipeline(cfg, store, http, lambda: llm, RunOptions()).run()

    assert stats["status"] == "ok", stats
    assert stats["jobs_seen"] == 10 and stats["jobs_new"] == 10
    assert stats["companies_failed"] == 1                         # Ashby 500 is logged, not fatal
    rows = {r["title"]: r for r in store.query("SELECT * FROM jobs")}
    assert rows["Senior Director, Commercial Strategy"]["filter_reason"].startswith("title excluded")
    assert rows["Strategy Analyst"]["filter_reason"] == "outside the US"
    assert rows["Commercial Analyst"]["filter_reason"].startswith("no sponsorship")
    assert rows["Market Access Analyst"]["filter_reason"].startswith("no sponsorship")   # found after hydration
    assert rows["Clinical Research Coordinator"]["filter_reason"] == "outside the US"
    assert rows["Research Associate, Downstream Process Development"]["band"] == "High"
    assert rows["Business Development Associate"]["sponsorship_flag"] == "positive"
    scored = [r for r in rows.values() if r["score"] is not None]
    assert len(scored) == stats["scored"] == len(llm.messages.calls) == 5
    assert store.scalar("SELECT COUNT(*) FROM llm_usage") == 5
    run = store.query("SELECT * FROM runs")[0]
    assert run["status"] == "ok" and run["cost_usd"] > 0

    # Workbook: tabs, rows, links, dropdown, formulas
    wb = load_workbook(cfg.path("workbook"))
    assert wb.sheetnames == ["Dashboard", "High Match", "Medium Match", "Low Match", "Applied", "Manual",
                             "Contacts", "Filtered Out"]
    high = wb["High Match"]
    header = [c.value for c in high[1]]
    assert header[-2:] == ["Status", "Notes"] and header[header.index("Status")] == "Status"
    titles = [high.cell(row=r, column=header.index("Role") + 1).value for r in range(2, high.max_row + 1)]
    assert "Research Associate, Downstream Process Development" in titles
    link = high.cell(row=2, column=header.index("Link") + 1)
    assert link.hyperlink is not None and link.hyperlink.target.startswith("https://")
    assert high.data_validations.dataValidation, "status dropdown missing"
    assert wb["Dashboard"]["B5"].value.startswith("=MAX(COUNTA('High Match'")
    assert wb["Filtered Out"].max_row >= 4

    # Riya marks a job "skip" and adds a note on another; the next rebuild keeps both.
    status_col = header.index("Status") + 1
    skip_id = high.cell(row=2, column=1).value
    high.cell(row=2, column=status_col, value="skip")
    if high.max_row >= 3:
        high.cell(row=3, column=status_col + 1, value="ask Mickey for intro")
    wb.save(cfg.path("workbook"))
    async with HttpClient(max_retries=0) as http:
        mock_boards()
        stats2 = await Pipeline(cfg, store, http, lambda: llm, RunOptions(fetch=False)).run()
    assert stats2["status"] == "ok"
    assert store.scalar("SELECT user_status FROM jobs WHERE id=?", (skip_id,)) == "skip"
    wb2 = load_workbook(cfg.path("workbook"))
    ids = [wb2["High Match"].cell(row=r, column=1).value for r in range(2, wb2["High Match"].max_row + 1)]
    assert skip_id not in ids                                    # skipped jobs leave the tabs
    assert len(llm.messages.calls) == 5                          # nothing re-scored on the second run

    # Logs: per-run file, JSON lines with run id, errors file has the Ashby failure
    logs = cfg.path("logs_dir")
    assert len(list((logs / "runs").glob("*.log"))) >= 1
    lines = [json.loads(l) for l in (logs / "jobagent.jsonl").read_text().splitlines()]
    assert {l["run_id"] for l in lines} - {"-"} and any(l["stage"] == "score" for l in lines)
    assert "AI Bio" in (logs / "errors.log").read_text()
    store.close()


@respx.mock
async def test_batch_mode(cfg):
    mock_boards()
    cfg.data["scoring"]["mode"] = "batch"
    cfg.data["scoring"]["batch_poll_seconds"] = 0
    store = Store(cfg.path("db"))
    llm = FakeLLM()
    async with HttpClient(max_retries=0) as http:
        stats = await Pipeline(cfg, store, http, lambda: llm, RunOptions()).run()
    assert stats["scored"] == 5 and store.scalar("SELECT status FROM pending_batches") == "collected"
    assert store.scalar("SELECT SUM(batch) FROM llm_usage") == 5
    store.close()


def test_no_api_key_still_builds_workbook(cfg):
    store = Store(cfg.path("db"))
    path = build_workbook(store, cfg.path("workbook"), cfg.data["workbook"])
    assert path.exists()
    store.close()


def test_run_lock(tmp_path):
    lock = tmp_path / "run.lock"
    with RunLock(lock, 6):
        with pytest.raises(RuntimeError):
            with RunLock(lock, 6):
                pass
    assert not lock.exists()


def test_secrets_masked():
    assert "sk-ant-***" in mask_secrets("key=sk-ant-api03-abcdefghijklmnop")
    assert mask_secrets("x-api-key: abc123") == "x-api-key: ***"
