"""Builds output/riya_job_tracker.xlsx from the database.

- Riya's edits in the Status and Notes columns are read back into the database before every rebuild,
  so nothing she types is lost.
- The file is written to a temp file and swapped in atomically. If it is open in Excel (locked on
  Windows), a timestamped copy is written instead and a warning is logged.
"""

from __future__ import annotations

import json
import logging
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..db import Store

log = logging.getLogger("jobagent.export")

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
EDIT_FILL = PatternFill("solid", start_color="FFF7D6")   # cells Riya edits
NEW_FILL = PatternFill("solid", start_color="E3F1E6")
THIN = Side(style="thin", color="D0D5DD")
AUTO_APPLY_ATS = {"greenhouse", "lever", "ashby"}

MATCH_COLUMNS = [  # (header, width, key)
    ("Job ID", 8, "id"), ("New", 6, "new"), ("Track", 10, "track"), ("Score", 7, "score"),
    ("Company", 22, "company"), ("Role", 38, "title"), ("Location", 24, "location"), ("Remote", 8, "remote"),
    ("Posted", 11, "posted_at"), ("Why", 60, "reason"), ("Matched skills", 36, "matched_skills"),
    ("Gaps", 30, "gaps"), ("Level asked", 11, "seniority_fit"), ("Sponsorship", 13, "sponsorship_flag"),
    ("Salary", 20, "salary"), ("Board", 14, "ats"), ("Apply method", 14, "apply_method"),
    ("Link", 12, "url"), ("Status", 16, "user_status"), ("Notes", 30, "user_note"),
]
APPLIED_COLUMNS = [("Job ID", 8), ("Track", 10), ("Company", 22), ("Role", 38), ("Applied on", 12),
                   ("Method", 12), ("Confirmation", 30), ("Link", 12), ("Status", 16), ("Notes", 30)]
MANUAL_COLUMNS = [("Job ID", 8), ("Track", 10), ("Score", 7), ("Company", 22), ("Role", 38), ("Why manual", 30),
                  ("Prepared answers", 40), ("Link", 12), ("Status", 16), ("Notes", 30)]
CONTACT_COLUMNS = [("Job ID", 8), ("Company", 22), ("Role", 38), ("Contact", 22), ("Title", 26), ("Email", 28),
                   ("Source", 12), ("Email status", 16), ("Last sent", 12), ("Next follow-up", 14), ("Reply", 10)]
FILTERED_COLUMNS = [("Company", 22), ("Role", 40), ("Location", 26), ("Reason", 50), ("Link", 12)]
EDITABLE_SHEETS = ("High Match", "Medium Match", "Low Match", "Applied", "Manual")


# ---- reading Riya's edits back ---------------------------------------------
def read_user_edits(path: Path) -> list[tuple[str, str, int]]:
    """Return (status, notes, job_id) for every row in the editable tabs."""
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corrupt / partially written file must never stop a run
        log.warning("could not read existing workbook %s: %s", path.name, exc)
        return []
    edits: dict[int, tuple[str, str]] = {}
    try:
        for name in EDITABLE_SHEETS:
            if name not in wb.sheetnames:
                continue
            rows = wb[name].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            idx = {h: i for i, h in enumerate(header) if h}
            if not {"Job ID", "Status", "Notes"} <= idx.keys():
                continue
            for row in rows:
                job_id = row[idx["Job ID"]]
                if not isinstance(job_id, (int, float)):
                    continue
                status = (row[idx["Status"]] or "").strip() if isinstance(row[idx["Status"]], str) else ""
                notes = (row[idx["Notes"]] or "").strip() if isinstance(row[idx["Notes"]], str) else ""
                edits[int(job_id)] = (status, notes)
    finally:
        wb.close()
    return [(s, n, jid) for jid, (s, n) in edits.items()]


# ---- writing ---------------------------------------------------------------
def _style_header(ws, columns) -> None:
    ws.append([c[0] for c in columns])
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col[1]
    ws.row_dimensions[1].height = 28
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.freeze_panes = "E2" if columns is MATCH_COLUMNS else "A2"


def _finish_table(ws, columns, n_rows: int, status_options: list[str]) -> None:
    last_col = get_column_letter(len(columns))
    last_row = max(2, n_rows + 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    headers = [c[0] for c in columns]
    body_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="1155CC", underline="single")
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for cell in row:
            cell.font = link_font if cell.hyperlink else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=headers[cell.column - 1] in (
                "Why", "Matched skills", "Gaps", "Role", "Notes", "Reason", "Prepared answers"))
            cell.border = Border(bottom=THIN)
    for name in ("Status", "Notes"):
        if name in headers:
            col = get_column_letter(headers.index(name) + 1)
            for r in range(2, n_rows + 2):
                ws[f"{col}{r}"].fill = EDIT_FILL
    if "Status" in headers:
        col = get_column_letter(headers.index("Status") + 1)
        options = ",".join(o for o in status_options if o)
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True, showDropDown=False)
        dv.error = "Pick a status from the list"
        dv.add(f"{col}2:{col}{max(last_row, 500)}")
        ws.add_data_validation(dv)
    if "Score" in headers and n_rows:
        col = get_column_letter(headers.index("Score") + 1)
        ws.conditional_formatting.add(f"{col}2:{col}{n_rows + 1}", ColorScaleRule(
            start_type="num", start_value=30, start_color="F8D7DA", mid_type="num", mid_value=60,
            mid_color="FFF3CD", end_type="num", end_value=90, end_color="C8E6C9"))


def _match_row(r: dict[str, Any], today: str) -> list[Any]:
    def joined(raw):
        try:
            return "; ".join(json.loads(raw or "[]"))
        except (TypeError, ValueError):
            return raw or ""
    method = "auto (phase 2)" if r["ats"] in AUTO_APPLY_ATS else "manual"
    values = {
        "id": r["id"], "new": "NEW" if (r["first_seen_at"] or "")[:10] == today else "",
        "track": (r["track"] or "").title(), "score": r["score"], "company": r["company"], "title": r["title"],
        "location": r["location"], "remote": "Yes" if r["is_remote"] else "", "posted_at": r["posted_at"] or "",
        "reason": r["reason"] or "", "matched_skills": joined(r["matched_skills"]), "gaps": joined(r["gaps"]),
        "seniority_fit": r["seniority_fit"] or "", "sponsorship_flag": (r["sponsorship_flag"] or "").replace("_", " "),
        "salary": r["salary"] or "", "ats": r["ats"], "apply_method": method, "url": "Open",
        "user_status": r["user_status"] or "", "user_note": r["user_note"] or "",
    }
    return [values[c[2]] for c in MATCH_COLUMNS]


def _write_match_sheet(wb: Workbook, title: str, rows: list, today: str, status_options: list[str]) -> None:
    ws = wb.create_sheet(title)
    _style_header(ws, MATCH_COLUMNS)
    link_col = [c[0] for c in MATCH_COLUMNS].index("Link") + 1
    new_col = [c[0] for c in MATCH_COLUMNS].index("New") + 1
    for i, r in enumerate(rows, start=2):
        ws.append(_match_row(r, today))
        url = r["apply_url"] or r["url"]
        if url:
            ws.cell(row=i, column=link_col).hyperlink = url
        if ws.cell(row=i, column=new_col).value:
            ws.cell(row=i, column=new_col).fill = NEW_FILL
    _finish_table(ws, MATCH_COLUMNS, len(rows), status_options)


def _write_simple_sheet(wb: Workbook, title: str, columns, rows: list[list[Any]], status_options,
                        links: list[str | None] | None = None) -> None:
    ws = wb.create_sheet(title)
    _style_header(ws, columns)
    headers = [c[0] for c in columns]
    for i, row in enumerate(rows, start=2):
        ws.append(row)
        if links and "Link" in headers and links[i - 2]:
            ws.cell(row=i, column=headers.index("Link") + 1).hyperlink = links[i - 2]
    _finish_table(ws, columns, len(rows), status_options)


def _write_dashboard(wb: Workbook, store: Store, stats: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    title_font = Font(name=FONT, bold=True, size=16, color="1F3A5F")
    label_font = Font(name=FONT, size=10, color="475467")
    value_font = Font(name=FONT, bold=True, size=12)
    small = Font(name=FONT, size=9, color="667085", italic=True)
    ws["A1"] = "Riya's job search tracker"
    ws["A1"].font = title_font
    ws["A2"] = f"Last updated {datetime.now().strftime('%a %b %d, %Y %I:%M %p')}"
    ws["A2"].font = label_font
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 15

    # Live counts (formulas, so they update as Riya edits statuses).
    metrics = [
        ("High matches", "=MAX(COUNTA('High Match'!A:A)-1,0)"),
        ("Medium matches", "=MAX(COUNTA('Medium Match'!A:A)-1,0)"),
        ("Low matches", "=MAX(COUNTA('Low Match'!A:A)-1,0)"),
        ("New today (High + Medium)", "=COUNTIF('High Match'!B2:B5000,\"NEW\")+COUNTIF('Medium Match'!B2:B5000,\"NEW\")"),
        ("Marked \"apply\"", "=COUNTIF('High Match'!S2:S5000,\"apply\")+COUNTIF('Medium Match'!S2:S5000,\"apply\")"
                            "+COUNTIF('Low Match'!S2:S5000,\"apply\")"),
        ("Marked \"skip\"", "=COUNTIF('High Match'!S2:S5000,\"skip\")+COUNTIF('Medium Match'!S2:S5000,\"skip\")"
                           "+COUNTIF('Low Match'!S2:S5000,\"skip\")"),
        ("Applied", "=MAX(COUNTA(Applied!A:A)-1,0)"),
        ("Waiting for manual apply", "=MAX(COUNTA(Manual!A:A)-1,0)"),
        ("Interviewing", "=COUNTIF('High Match'!S2:S5000,\"interviewing\")+COUNTIF('Medium Match'!S2:S5000,\"interviewing\")"
                         "+COUNTIF(Applied!I2:I5000,\"interviewing\")"),
    ]
    ws["A4"] = "Pipeline"
    ws["A4"].font = Font(name=FONT, bold=True, size=11)
    for i, (label, formula) in enumerate(metrics, start=5):
        ws[f"A{i}"] = label
        ws[f"A{i}"].font = label_font
        ws[f"B{i}"] = formula
        ws[f"B{i}"].font = value_font
    row = 5 + len(metrics) + 1

    by_track = store.query("SELECT COALESCE(track,'?') t, band, COUNT(*) n FROM jobs WHERE is_active=1 "
                           "AND band IS NOT NULL GROUP BY t, band")
    ws[f"A{row}"] = "Matches by track"
    ws[f"A{row}"].font = Font(name=FONT, bold=True, size=11)
    row += 1
    for c, h in enumerate(["Track", "High", "Medium", "Low"], start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.border = Border(bottom=THIN)
    counts = {(r["t"], r["band"]): r["n"] for r in by_track}
    for track in ("business", "research"):
        row += 1
        ws.cell(row=row, column=1, value=track.title()).font = label_font
        for c, band in enumerate(("High", "Medium", "Low"), start=2):
            ws.cell(row=row, column=c, value=counts.get((track, band), 0)).font = Font(name=FONT, size=10)
    row += 2

    ws[f"A{row}"] = "Recent runs"
    ws[f"A{row}"].font = Font(name=FONT, bold=True, size=11)
    row += 1
    run_headers = ["Run", "Status", "Companies", "Jobs seen", "New jobs", "Scored", "AI cost ($)"]
    for c, h in enumerate(run_headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.border = Border(bottom=THIN)
    for r in store.query("SELECT * FROM runs ORDER BY started_at DESC LIMIT 14"):
        row += 1
        s = json.loads(r["stats"] or "{}")
        vals = [r["run_id"], r["status"], s.get("companies_ok"), s.get("jobs_seen"), s.get("jobs_new"),
                s.get("scored"), round(r["cost_usd"] or 0, 3)]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            if c == 7:
                cell.number_format = "$#,##0.000"
    row += 2
    month_cost = store.spend_since(datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0).isoformat())
    ws[f"A{row}"] = "AI spend this month ($)"
    ws[f"A{row}"].font = label_font
    ws[f"B{row}"] = round(month_cost, 2)
    ws[f"B{row}"].number_format = "$#,##0.00"
    ws[f"B{row}"].font = value_font
    row += 2
    ws[f"A{row}"] = ("How to use: edit only the yellow Status and Notes columns. Set Status to \"skip\" to hide a job, "
                     "or \"apply\" to queue a Low match. Everything else is rebuilt each morning.")
    ws[f"A{row}"].font = small
    ws[f"A{row + 1}"] = ("Score: 0-100 fit from Claude. High >= 75, Medium 50-74. \"NEW\" = first seen today. "
                         "Apply method \"auto\" = form can be auto-filled once phase 2 is on.")
    ws[f"A{row + 1}"].font = small


def build_workbook(store: Store, path: Path, cfg: dict, stats: dict[str, Any] | None = None) -> Path:
    """Write the workbook; returns the path actually written."""
    status_options = list(cfg.get("status_options", []))
    today = datetime.now(timezone.utc).date().isoformat()
    hidden = "('skip','rejected')"
    base = (f"SELECT * FROM jobs WHERE is_active=1 AND band=? AND COALESCE(user_status,'') NOT IN {hidden} "
            "AND COALESCE(apply_status,'') NOT IN ('applied','manual') ORDER BY score DESC, first_seen_at DESC")
    wb = Workbook()
    _write_dashboard(wb, store, stats or {})
    _write_match_sheet(wb, "High Match", store.query(base, ("High",)), today, status_options)
    _write_match_sheet(wb, "Medium Match", store.query(base, ("Medium",)), today, status_options)
    low_limit = int(cfg.get("low_match_max_rows", 400))
    _write_match_sheet(wb, "Low Match", store.query(base + f" LIMIT {low_limit}", ("Low",)), today, status_options)

    applied = store.query("SELECT * FROM jobs WHERE apply_status='applied' ORDER BY applied_at DESC")
    _write_simple_sheet(wb, "Applied", APPLIED_COLUMNS,
                        [[r["id"], (r["track"] or "").title(), r["company"], r["title"], (r["applied_at"] or "")[:10],
                          r["apply_method"] or "", "", "Open", r["user_status"] or "", r["user_note"] or ""]
                         for r in applied], status_options, [r["apply_url"] or r["url"] for r in applied])
    manual = store.query("SELECT * FROM jobs WHERE apply_status='manual' ORDER BY score DESC")
    _write_simple_sheet(wb, "Manual", MANUAL_COLUMNS,
                        [[r["id"], (r["track"] or "").title(), r["score"], r["company"], r["title"],
                          r["apply_error"] or "", "", "Open", r["user_status"] or "", r["user_note"] or ""]
                         for r in manual], status_options, [r["apply_url"] or r["url"] for r in manual])
    _write_simple_sheet(wb, "Contacts", CONTACT_COLUMNS, [], status_options)

    since = (datetime.now(timezone.utc) - timedelta(days=int(cfg.get("filtered_out_days", 3)))).isoformat()
    filtered = store.query("SELECT company, title, location, filter_reason, url FROM jobs WHERE "
                           "filter_status='excluded' AND filtered_at>=? AND filter_reason NOT LIKE 'title not relevant%' "
                           "ORDER BY filtered_at DESC LIMIT 1000", (since,))
    _write_simple_sheet(wb, "Filtered Out", FILTERED_COLUMNS,
                        [[r["company"], r["title"], r["location"], r["filter_reason"], "Open"] for r in filtered],
                        status_options, [r["url"] for r in filtered])

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem} ({datetime.now().strftime('%Y%m%d-%H%M')}).xlsx")
        os.replace(tmp, alt)
        log.warning("%s is open (probably in Excel); wrote %s instead. Close Excel before the next run.",
                    path.name, alt.name)
        return alt
